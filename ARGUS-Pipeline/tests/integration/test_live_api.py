"""Live contract tests for the seeded Compose stack.

Run only against a disposable synthetic-data environment:
    ARGUS_LIVE_TESTS=1 pytest -q tests/integration/test_live_api.py
"""

from __future__ import annotations

import os

import pytest
from fastapi.testclient import TestClient

if not os.getenv("ARGUS_LIVE_TESTS"):
    pytest.skip("Set ARGUS_LIVE_TESTS=1 to run against the live synthetic stack.", allow_module_level=True)

from main import app


@pytest.fixture(scope="module")
def client() -> TestClient:
    return TestClient(app)


def login(client: TestClient, employee_id: str, password: str) -> str:
    response = client.post("/api/v1/auth/login", json={"employee_id": employee_id, "password": password})
    assert response.status_code == 200, response.text
    return response.json()["access_token"]


def test_auth_success_and_failure(client: TestClient) -> None:
    token = login(client, "admin@demo.com", "password123")
    assert token
    assert client.post(
        "/api/v1/auth/login",
        json={"employee_id": "admin@demo.com", "password": "wrong-password"},
    ).status_code == 401


def test_sensitive_case_requires_justification(client: TestClient) -> None:
    token = login(client, "admin@demo.com", "password123")
    headers = {"Authorization": f"Bearer {token}"}
    cases = client.get("/api/v1/cases", headers=headers)
    assert cases.status_code == 200
    case = next(item for item in cases.json() if item["is_sensitive"])
    denied = client.get(f"/api/v1/cases/{case['case_id']}", headers=headers)
    assert denied.status_code == 428
    allowed = client.get(
        f"/api/v1/cases/{case['case_id']}?justification=Subject%20is%20a%20Politically%20Exposed%20Person",
        headers=headers,
    )
    assert allowed.status_code == 200
    audit = client.get("/api/v1/audit", headers=headers)
    assert audit.status_code == 200
    assert any(
        entry["action"] == "SENSITIVE_VIEW"
        and entry["resource"] == f"case:{case['case_id']}"
        and entry["justification"] == "Subject is a Politically Exposed Person"
        for entry in audit.json()
    )


def test_ingestion_health_is_live_and_role_protected(client: TestClient) -> None:
    admin = login(client, "admin@demo.com", "password123")
    response = client.get("/api/v1/admin/ingestion-health", headers={"Authorization": f"Bearer {admin}"})
    assert response.status_code == 200
    assert {"status", "lastRun", "processed"} <= response.json().keys()

    investigator = login(client, "INV001", "demo123")
    denied = client.get("/api/v1/admin/ingestion-health", headers={"Authorization": f"Bearer {investigator}"})
    assert denied.status_code == 403


def test_graph_alerts_and_export_are_live(client: TestClient) -> None:
    token = login(client, "admin@demo.com", "password123")
    headers = {"Authorization": f"Bearer {token}"}
    graph = client.get("/api/v1/network/accused", params={"accused_name": "Alpha Holdings"}, headers=headers)
    assert graph.status_code == 200
    payload = graph.json()
    if "phone_contacts" not in payload:
        pytest.skip("Alpha Holdings AML narrative not loaded; run scripts/seed_demo.py on a disposable stack.")
    assert len(payload["phone_contacts"]) >= 4

    alerts = client.get("/api/v1/alerts", headers=headers)
    assert alerts.status_code == 200
    assert any(alert["entity_value"] == "Alpha Holdings" for alert in alerts.json())

    case = client.get("/api/v1/cases", headers=headers).json()[0]
    export = client.post(
        "/api/v1/reports/export",
        headers=headers,
        json={
            "case_id": case["case_id"],
            "sections": ["Entity List", "Notes"],
            "justification": "Live export verification for authorized investigation",
        },
    )
    assert export.status_code == 200
    assert export.headers["content-type"].startswith("application/pdf")
    assert "attachment" in export.headers["content-disposition"]


# ---------------------------------------------------------------------------
# Jurisdiction scoping (PRD FR-10). The seeded investigator INV001 belongs to "Central District"; admin is unscoped.
# Needs the demo data loaded and `scripts/backfill_jurisdiction.py` run so every FIR carries a jurisdiction.
# ---------------------------------------------------------------------------

DISTRICT = "Central District"


def _auth(client: TestClient, employee_id: str, password: str) -> dict:
    return {"Authorization": f"Bearer {login(client, employee_id, password)}"}


@pytest.fixture(scope="module")
def investigator(client: TestClient) -> dict:
    return _auth(client, "INV001", "demo123")


@pytest.fixture(scope="module")
def admin(client: TestClient) -> dict:
    return _auth(client, "admin@demo.com", "password123")


def _fir_jurisdictions() -> dict:
    """fir_id -> (jurisdiction, accused) straight from Elasticsearch."""
    from elasticsearch import helpers
    from main import es_client

    out = {}
    for hit in helpers.scan(es_client, index="argus-firs", query={"query": {"exists": {"field": "accused"}}},
                            _source=["fir_id", "accused", "jurisdiction"]):
        source = hit["_source"]
        out[source["fir_id"]] = (source.get("jurisdiction"), source["accused"])
    return out


def test_search_results_only_contain_the_investigators_jurisdiction(client, investigator, admin) -> None:
    scoped = client.get("/api/v1/search/firs", params={"q": "Police"}, headers=investigator).json()
    everything = client.get("/api/v1/search/firs", params={"q": "Police"}, headers=admin).json()
    assert scoped["count"] > 0
    assert {r["jurisdiction"] for r in scoped["results"]} == {DISTRICT}
    # The unscoped role sees other districts too, so the filter (not the data) is what limits the investigator.
    assert len({r["jurisdiction"] for r in everything["results"]}) > 1


def test_dossier_and_judges_view_are_scoped(client, investigator, admin) -> None:
    firs = _fir_jurisdictions()
    elsewhere = next(fid for fid, (j, _) in firs.items() if j not in (DISTRICT, None))
    own = next(fid for fid, (j, _) in firs.items() if j == DISTRICT)

    hidden = client.get("/api/v1/search/master-dossier", params={"query": elsewhere}, headers=investigator).json()
    assert elsewhere not in {r["fir_id"] for r in hidden["dossier_records"]}
    assert all(r["jurisdiction"] == DISTRICT for r in hidden["dossier_records"])
    visible = client.get("/api/v1/search/master-dossier", params={"query": own}, headers=investigator).json()
    assert own in {r["fir_id"] for r in visible["dossier_records"]}
    assert client.get("/api/v1/search/master-dossier", params={"query": elsewhere}, headers=admin).json()["total_matches"] == 1

    judges = client.get("/api/v1/search/judges-view", headers=investigator).json()
    assert judges["count"] > 0 and all(r["jurisdiction"] == DISTRICT for r in judges["results"])


def _graph_people() -> dict:
    """suspect -> {fir_id: jurisdiction}, from the same Neo4j graph the endpoint reads."""
    from main import neo4j_driver

    people: dict = {}
    with neo4j_driver.session() as session:
        for row in session.run("MATCH (s:Suspect)-[:LINKED_TO_FIR]->(f:FIR) RETURN s.id AS s, f.id AS f, f.jurisdiction AS j"):
            people.setdefault(row["s"], {})[row["f"]] = row["j"]
    return people


def test_out_of_scope_accused_is_indistinguishable_from_nonexistent(client, investigator, admin) -> None:
    # In the demo graph everyone has a FIR in every district, so create a person who exists only in another one.
    from main import neo4j_driver

    name = "ZZ Scope Test Person"
    with neo4j_driver.session() as session:
        session.run("MERGE (s:Suspect {id: $n}) MERGE (f:FIR {id: 'ZZ-SCOPE-1'}) SET f.jurisdiction = 'Harbor Division' "
                    "MERGE (s)-[:LINKED_TO_FIR]->(f)", n=name)
    try:
        hidden = client.get("/api/v1/network/accused", params={"accused_name": name}, headers=investigator)
        missing = client.get("/api/v1/network/accused", params={"accused_name": "Nobody At All Exists"}, headers=investigator)
        assert hidden.status_code == missing.status_code == 200          # never a 403 that would confirm it exists
        assert hidden.json() == missing.json() == {"status": "success", "message": "No results in your jurisdiction."}
        assert client.get("/api/v1/network/accused", params={"accused_name": name}, headers=admin).json()["total_firs"] == 1
    finally:
        with neo4j_driver.session() as session:
            session.run("MATCH (n) WHERE n.id IN [$n, 'ZZ-SCOPE-1'] DETACH DELETE n", n=name)


def test_person_in_several_districts_only_shows_in_scope_firs(client, investigator, admin) -> None:
    people = _graph_people()
    name = next(p for p, m in people.items() if DISTRICT in m.values() and len(set(m.values())) > 1)
    scoped = client.get("/api/v1/network/accused", params={"accused_name": name}, headers=investigator).json()
    full = client.get("/api/v1/network/accused", params={"accused_name": name}, headers=admin).json()
    assert 0 < scoped["total_firs"] < full["total_firs"]
    assert all(people[name][fid] == DISTRICT for fid in scoped["linked_firs"])


def test_applied_scope_is_recorded_in_the_audit_log(client, investigator, admin) -> None:
    client.get("/api/v1/search/firs", params={"q": "Police"}, headers=investigator)
    entries = client.get("/api/v1/audit", params={"action": "search_firs"}, headers=admin).json()
    assert any((e.get("details") or {}).get("jurisdiction_filter") == DISTRICT for e in entries)
    assert client.get("/api/v1/audit/verify", headers=admin).json()["valid"] is True


def test_women_safety_patterns_are_present_evidence_based_and_scoped(client, investigator, admin) -> None:
    rows = client.get("/api/v1/patterns", headers=admin).json()
    by_type = {}
    for p in rows:
        by_type.setdefault(p["pattern_type"], []).append(p)
    assert {"phone_cluster_hub", "repeat_offender_recurrence", "co_accused_cluster"} <= set(by_type)
    assert "burner_phone_cluster" not in by_type                       # renamed
    required = {"pattern_id", "pattern_type", "confidence", "description", "explanation", "entities", "detected_at", "status", "source"}
    assert all(required <= set(p) for p in rows)                       # the UI reads these keys
    flagged = [p for p in rows if p.get("women_safety_flag")]
    assert flagged and len(flagged) < len(rows)                        # meaningful, not blanket
    assert all(p["women_safety_fir_count"] >= 2 and "Women-safety relevance" in p["explanation"] for p in flagged)
    assert all(p["risk_tier"] in {"HIGH", "MEDIUM"} for p in by_type["repeat_offender_recurrence"] + by_type["co_accused_cluster"])

    # A scoped investigator's recurrence counts only cover their own jurisdiction.
    scoped = {p["pattern_id"]: p for p in client.get("/api/v1/patterns", headers=investigator).json()}
    full = {p["pattern_id"]: p for p in rows}
    for pid, p in scoped.items():
        if p["pattern_type"] == "repeat_offender_recurrence" and pid in full:
            assert int(p["description"].split(" appears in ")[1].split(" FIRs")[0]) <= int(full[pid]["description"].split(" appears in ")[1].split(" FIRs")[0])


# ---------------------------------------------------------------------------
# DOCX export
# ---------------------------------------------------------------------------

def _export(client, headers, case_id, fmt=None, **body):
    params = {"format": fmt} if fmt else {}
    payload = {"case_id": case_id, "sections": ["Entity List", "Notes", "Pattern Findings", "Source Citations"], **body}
    return client.post("/api/v1/reports/export", params=params, json=payload, headers=headers)


def test_docx_export_end_to_end(client, admin) -> None:
    from io import BytesIO

    from docx import Document

    cases = client.get("/api/v1/cases", headers=admin).json()
    plain = next(c for c in cases if not c["is_sensitive"])
    response = _export(client, admin, plain["case_id"], "docx", justification="Court submission")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert response.headers["content-disposition"] == f"attachment; filename=argus-report-{plain['case_id']}.docx"
    text = "\n".join(p.text for p in Document(BytesIO(response.content)).paragraphs)
    assert plain["title"] in text
    assert "Exported by Demo Administrator at " in text and "justification: Court submission" in text
    # Patterns in the report are the case's own, each with the AI caveat.
    assert text.count("AI-derived lead — verify before use") >= text.count("% confidence")

    # PDF is still the default, and unknown formats are refused.
    assert _export(client, admin, plain["case_id"]).content[:5] == b"%PDF-"
    assert _export(client, admin, plain["case_id"], "rtf").status_code == 422


def test_docx_export_respects_sensitive_case_gate_and_is_audited(client, admin) -> None:
    sensitive = next(c for c in client.get("/api/v1/cases", headers=admin).json() if c["is_sensitive"])
    assert _export(client, admin, sensitive["case_id"], "docx").status_code == 428
    assert _export(client, admin, sensitive["case_id"], "docx", justification="Warrant 42").status_code == 200
    entries = client.get("/api/v1/audit", params={"action": "export_report"}, headers=admin).json()
    assert any((e.get("details") or {}).get("format") == "docx" and e["resource"].endswith(".docx") for e in entries)


def test_report_only_lists_patterns_involving_the_cases_own_entities(client, admin) -> None:
    from io import BytesIO

    from docx import Document

    all_patterns = client.get("/api/v1/patterns", headers=admin).json()
    chosen = None
    for case in client.get("/api/v1/cases", headers=admin).json():
        detail = client.get(f"/api/v1/cases/{case['case_id']}", params={"justification": "test"}, headers=admin).json()
        pinned = {e["entity_value"] for e in detail["entities"]}
        expected = sum(1 for p in all_patterns if pinned & set(p["entities"]))
        if expected:
            chosen = (case, expected)
            break
    assert chosen, "the demo data should have a case with matching patterns"
    case, expected = chosen
    response = _export(client, admin, case["case_id"], "docx", justification="test")
    text = "\n".join(p.text for p in Document(BytesIO(response.content)).paragraphs)
    assert text.count("% confidence") == expected < len(all_patterns)      # not all 71 patterns in the system
