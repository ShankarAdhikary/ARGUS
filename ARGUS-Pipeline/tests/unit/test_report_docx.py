"""The .docx report: content, per-finding AI caveat, and the audit footer."""

from datetime import datetime, timezone
from io import BytesIO

from docx import Document

from report_docx import AI_LEAD_NOTICE, build_case_docx

CASE = {
    "case_id": "c-1", "title": "Trafficking Ring - Central", "fir_number": "FIR-2026-0018", "jurisdiction": "Central District",
    "status": "open", "is_sensitive": True, "sensitivity_reason": "Involves a minor", "category": "trafficking",
    "opened_at": "2026-09-09T10:00:00+00:00",
    "entities": [{"entity_type": "person", "entity_value": "Vikram Singh", "linked_by": "Rahul Verma", "linked_at": "x"}],
    "notes": [{"note_id": "n", "author": "Ayesha Khan", "content": "Met informant", "created_at": "x"}],
}
PATTERNS = [
    {"description": "Vikram Singh appears in 13 FIRs.", "confidence": 0.95, "explanation": "Linked to 13 FIRs.", "women_safety_flag": True},
    {"description": "Phone 9999988888 is a hub.", "confidence": 0.71, "explanation": ""},
]


def _read(**kw):
    args = dict(case=CASE, sections=["Entity List", "Notes", "Pattern Findings", "Source Citations"], patterns=PATTERNS,
                exported_by="Rahul Verma", exported_at=datetime(2026, 9, 29, 12, 30, tzinfo=timezone.utc), justification="Court submission")
    args.update(kw)
    doc = Document(BytesIO(build_case_docx(args["case"], args["sections"], args["patterns"], args["exported_by"], args["exported_at"], args["justification"])))
    return doc, "\n".join(p.text for p in doc.paragraphs), [c.text for t in doc.tables for r in t.rows for c in r.cells]


def test_contains_every_selected_section_and_metadata():
    doc, text, cells = _read()
    assert "Trafficking Ring - Central" in text
    assert {"FIR-2026-0018", "Central District", "Yes — Involves a minor"} <= set(cells)
    assert "Vikram Singh" in cells and "Met informant" in text and "FIR-2026-0018: Vikram Singh" in text
    assert "95% confidence" in text and "71% confidence" in text and "[Women Safety]" in text


def test_every_finding_carries_the_ai_caveat():
    _, text, _ = _read()
    assert text.count(AI_LEAD_NOTICE) == len(PATTERNS)


def test_audit_footer_names_who_when_and_why():
    _, text, _ = _read()
    assert "Exported by Rahul Verma at 2026-09-29 12:30:00 UTC, justification: Court submission" in text


def test_unselected_sections_are_left_out():
    _, text, cells = _read(sections=["Notes"])
    assert "Met informant" in text and "Pattern Findings" not in text and "Vikram Singh" not in cells


def test_no_patterns_message():
    _, text, _ = _read(patterns=[])
    assert "No detected patterns involve this case's pinned entities." in text
